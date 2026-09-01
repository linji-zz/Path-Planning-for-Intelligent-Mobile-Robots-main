"""md_to_docx.py — Convert paper_draft.md to paper_draft.docx (Chinese paper).

Handles headings (#/##/###/####), markdown tables, bullet lists, inline bold,
and embeds the referenced figures (fig_*.png) at their citation points with
captions. Body = 宋体, headings = 黑体, A4 page.
"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import parse_xml

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "paper_draft.md")
OUT = os.path.join(BASE, "paper_draft.docx")

FIG_CAPTIONS = {
    "fig_rl_interaction.png": "图1 强化学习交互过程",
    "fig_dqn_structure.png": "图2 DQN训练结构",
    "fig_cat_structure.png": "图3 CAT-DQN网络结构",
    "fig_lambda_curve.png": "图4 ASGS修正强度λ随训练进度变化",
    "fig_training_flow.png": "图5 完整方法训练流程",
    "fig_grid_maps.png": "图6 栅格地图",
    "fig_paths_20x20_all.png": "图7 20×20地图各方法路径对比",
    "fig_reward_20x20_all.png": "图8 20×20地图奖励收敛曲线",
    "fig_paths_30x30_all.png": "图9 30×30地图各方法路径对比",
    "fig_reward_30x30_all.png": "图10 30×30地图奖励收敛曲线",
    "fig_trad_compare.png": "图11 与传统算法对比",
    "fig_noise_robustness.png": "图12 噪声鲁棒性对比",
}
FIG_RE = re.compile(r"fig_[A-Za-z0-9_]+\.png")

FIG_EN = {
    "fig_rl_interaction.png": "Fig.1 Reinforcement learning interaction process",
    "fig_dqn_structure.png": "Fig.2 DQN training structure",
    "fig_cat_structure.png": "Fig.3 CAT-DQN network structure",
    "fig_lambda_curve.png": "Fig.4 ASGS strength λ vs training progress",
    "fig_training_flow.png": "Fig.5 Training flow of the complete method",
    "fig_grid_maps.png": "Fig.6 Grid maps",
    "fig_paths_20x20_all.png": "Fig.7 Path comparison of all methods on the 20×20 map",
    "fig_reward_20x20_all.png": "Fig.8 Reward convergence curves on the 20×20 map",
    "fig_paths_30x30_all.png": "Fig.9 Path comparison of all methods on the 30×30 map",
    "fig_reward_30x30_all.png": "Fig.10 Reward convergence curves on the 30×30 map",
    "fig_trad_compare.png": "Fig.11 Comparison with traditional algorithms",
    "fig_noise_robustness.png": "Fig.12 Noise robustness comparison",
}

TABLE_EN = {
    "1": "Tab.1 Experimental results on the 20×20 map (mean±std, 3 seeds)",
    "2": "Tab.2 Experimental results on the 30×30 map (mean±std, 3 seeds)",
    "3": "Tab.3 Comparison with improved DQN algorithms (mean±std, 3 seeds)",
    "4": "Tab.4 Comparison with traditional algorithms",
    "5": "Tab.5 Test success rate of clean-trained models under sensor noise (mean±std, 5 seeds, 200 greedy episodes per seed)",
}


def set_font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def add_inline(paragraph, text, name="宋体", size=12):
    """Add text with **bold** inline formatting."""
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_font(run, name, size, bold=(i % 2 == 1))
    return paragraph


def add_figure(doc, filename):
    path = os.path.join(BASE, filename)
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cap.add_run(FIG_CAPTIONS.get(filename, filename)), "宋体", 10.5)
    en = FIG_EN.get(filename)
    if en:
        cap.add_run().add_break()
        set_font(cap.add_run(en), "Times New Roman", 10.5)


# ---------------- OMML (Word native equation) helpers ----------------
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _r(t):
    return f'<m:r><m:t>{t}</m:t></m:r>'


def _sub(e, s):
    return f'<m:sSub><m:e>{e}</m:e><m:sub>{s}</m:sub></m:sSub>'


def _sup(e, s):
    return f'<m:sSup><m:e>{e}</m:e><m:sup>{s}</m:sup></m:sSup>'


def _frac(n, d):
    return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'


def _paren(inner):
    return (f'<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
            f'<m:e>{inner}</m:e></m:d>')


def _sum(sub, e):
    return (f'<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{_r("")}</m:sup><m:e>{e}</m:e></m:nary>')


def _brace_arr(rows):
    eq = "".join(f'<m:e>{r}</m:e>' for r in rows)
    return (f'<m:d><m:dPr><m:begChr m:val="{{"/><m:endChr m:val=""/></m:dPr>'
            f'<m:e><m:eqArr>{eq}</m:eqArr></m:e></m:d>')


def _omath(inner):
    return (f'<m:oMathPara xmlns:m="{MATH_NS}"><m:oMath>{inner}</m:oMath></m:oMathPara>')


FORMULAS = {
    "R_total = r_dist + r_align + r_open": _omath(
        _sub(_r("R"), _r("total")) + _r("=")
        + _sub(_r("r"), _r("dist")) + _r("+")
        + _sub(_r("r"), _r("align")) + _r("+")
        + _sub(_r("r"), _r("open"))
    ),
    "r_dist = (d_{t-1} - d_t) * beta": _omath(
        _sub(_r("r"), _r("dist")) + _r("=")
        + _paren(_sub(_r("d"), _r("t-1")) + _r("−") + _sub(_r("d"), _r("t")))
        + _r("·") + _r("β")
    ),
    "r_align = alpha * cos(theta - theta_goal)": _omath(
        _sub(_r("r"), _r("align")) + _r("=") + _r("α") + _r("·") + _r("cos")
        + _paren(_r("θ") + _r("−") + _sub(_r("θ"), _r("goal")))
    ),
    "r_open = gamma * (1 - N_obs / 9)": _omath(
        _sub(_r("r"), _r("open")) + _r("=") + _r("γ") + _r("·")
        + _paren(_r("1") + _r("−") + _frac(_sub(_r("N"), _r("obs")), _r("9")))
    ),
    "P(a_i) = w_i / sum(w_j),  w_i = 0.1 if obs_i = 1 else 1.0": _omath(
        _r("P") + _paren(_sub(_r("a"), _r("i"))) + _r("=")
        + _sub(_r("w"), _r("i")) + _r("/")
        + _sum(_r("j"), _sub(_r("w"), _r("j"))) + _r(",  ")
        + _sub(_r("w"), _r("i")) + _r("=")
        + _brace_arr([
            _r("0.1, if ") + _sub(_r("obs"), _r("i")) + _r(" = 1"),
            _r("1.0, otherwise"),
        ])
    ),
    "Q'(s, a_i) = Q(s, a_i) - lambda * obs_i": _omath(
        _sup(_r("Q"), _r("′")) + _paren(_r("s,") + _sub(_r("a"), _r("i"))) + _r("=")
        + _r("Q") + _paren(_r("s,") + _sub(_r("a"), _r("i"))) + _r("−")
        + _r("λ") + _r("·") + _sub(_r("obs"), _r("i"))
    ),
    "lambda = 5.0 + 10.0 * (k / K)": _omath(
        _r("λ") + _r("=") + _r("5.0") + _r("+") + _r("10.0") + _r("·") + _frac(_r("k"), _r("K"))
    ),
}


def add_math(paragraph, omath_xml):
    paragraph._p.append(parse_xml(omath_xml))


def build_docx():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.0)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # blank line
        if not stripped:
            i += 1
            continue

        # equation (Word native OMML)
        if stripped in FORMULAS:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_math(p, FORMULAS[stripped])
            i += 1
            continue

        # table: consecutive lines starting with '|'
        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            # drop the separator row (|---|...)
            rows = [ln for ln in table_lines if not re.match(r"^\|[\s\-|:]+\|$", ln)]
            if not rows:
                continue
            header = [c.strip() for c in rows[0].strip("|").split("|")]
            ncols = len(header)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                cells = [c.strip() for c in row.strip("|").split("|")]
                # pad/trim to ncols
                cells = (cells + [""] * ncols)[:ncols]
                tr_cells = table.add_row().cells
                for c_idx, cell_text in enumerate(cells):
                    cell_obj = tr_cells[c_idx]
                    cell_obj.paragraphs[0].text = ""
                    add_inline(cell_obj.paragraphs[0], cell_text, size=10.5)
                    if r_idx == 0:
                        for run in cell_obj.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        # headings
        if stripped.startswith("#### "):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[5:]), "黑体", 12, bold=True)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[4:]), "黑体", 12, bold=True)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            set_font(p.add_run(stripped[3:]), "黑体", 14, bold=True)
        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(stripped[2:]), "黑体", 16, bold=True)
        # bullet list
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:], size=12)
        # figure-caption lines (图N ...) — keep as centered caption
        elif re.match(r"^图\d", stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, stripped, size=10.5)
        # table-caption lines (表N ...) — centered + bilingual
        elif re.match(r"^表(\d)", stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(stripped), "宋体", 10.5)
            m = re.match(r"^表(\d)", stripped)
            if m and m.group(1) in TABLE_EN:
                p.add_run().add_break()
                set_font(p.add_run(TABLE_EN[m.group(1)]), "Times New Roman", 10.5)
        # regular paragraph (may contain figure references)
        else:
            text = stripped
            figs = FIG_RE.findall(text)
            cleaned = re.sub(r"[（(][^）)]*fig_[^）)]*[）)]", "", text).strip()
            p = doc.add_paragraph()
            add_inline(p, cleaned, size=12)
            for fig in figs:
                add_figure(doc, fig)
        i += 1

    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build_docx()
